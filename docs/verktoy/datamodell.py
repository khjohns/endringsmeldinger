"""Lager tabellbeskrivelsen og hendelseskatalogen fra registrene i docs/datamodell/.

Kjøres fra repo-roten: /tmp/venv/bin/python docs/verktoy/datamodell.py
Skriver docs/datamodell/tabeller.md, hendelser.md, relasjoner.md og
tabellbeskrivelse.xlsx, men bare når innholdet er endret. Hendelseskatalogen
leser backend-modellene, og Excel-fila krever openpyxl, så begge trenger
venv-en i AGENTS.md. Katalogen over skjemaet, katalog.json, lages av
katalog.py mot en testbase.
"""

import importlib.util
import pathlib
import re
import sys

import tomllib


def _last(navn: str):
    sti = pathlib.Path(__file__).resolve().parent / f"{navn}.py"
    spesifikasjon = importlib.util.spec_from_file_location(navn, sti)
    modul = importlib.util.module_from_spec(spesifikasjon)
    spesifikasjon.loader.exec_module(modul)
    return modul


hk = _last("hendelseskatalog")
kat = _last("katalog")
rel = _last("relasjoner")
df = _last("dataflyt")

ROT = pathlib.Path(__file__).resolve().parents[2]
MAPPE = ROT / "docs" / "datamodell"
REGISTER = MAPPE / "tabeller.toml"
MARKDOWN = MAPPE / "tabeller.md"
EXCEL = MAPPE / "tabellbeskrivelse.xlsx"
RELASJONER = MAPPE / "relasjoner.md"
HOVEDPLAN = ROT / "docs" / "plans" / "2026-09-16-godkjenning-og-varig-levering.md"

IKT_KOLONNER = [
    "Tabellnavn",
    "Beskrivelse",
    "Er dette en transaksjonstabell?",
    "Hvor fra appen lagres og redigeres",
    "Er dette en grunndatatabell?",
    "Tabell der data hentes inn, f.eks. via Catenda-API-et",
    "Beskrivelse av relasjon til en annen tabell",
]
# Valgt av oppdragsgiver 29.09. Belegg, funn og skrivere står bare i tabeller.md.
TILLEGGSKOLONNER = ["Lagring", "Type", "Status", "Personopplysninger"]
UTENFOR_KOLONNER = ["Navn", "Lagring", "Beskrivelse", "Hvor fra appen lagres"]
UTENFOR_EXCEL = ("navn", "lagring", "beskrivelse", "lagres_fra")

TYPER = {
    "transaksjon": "Rader som registrerer noe en part eller bruker har gjort",
    "grunndata": "Prosjekter, brukere, medlemskap og oppsett som andre rader viser til",
    "avledet": "Kan bygges opp igjen fra andre tabeller og rettes ikke for hånd",
    "kopi": "Kopi av data fra Catenda",
    "teknisk": "Sesjoner, synkroniseringsstatus, køer og kvitteringer",
}
LAGRING = ("PostgreSQL", "SQLite")
STATUS = ("i bruk", "rest", "uavklart")
GJENOPPBYGGING = ("ja", "nei", "delvis", "ikke kontrollert", "ikke relevant")
BELEGG = {
    "K": "kjørt og observert (med dato)",
    "L": "lest i kode eller migrasjonsfil (med dato)",
    "D": "kontrollert i databasekatalogen tidligere (med dato)",
    "H": "historisk dokumentert i en audit",
}
TABELLFELT = (
    "navn",
    "lagring",
    "type",
    "status",
    "beskrivelse",
    "lagres_fra",
    "fra_catenda",
    "kan_gjenoppbygges",
    "personopplysninger",
    "skrivere",
    "funn",
    "belegg",
)
UTENFOR_FELT = ("navn", "lagring", "beskrivelse", "lagres_fra", "belegg")
FUNN_ID = re.compile(r"^[A-Z][A-Z0-9]*-\d{2}$")

MERKNAD = (
    "Generert fra docs/datamodell/tabeller.toml av docs/verktoy/datamodell.py. "
    "Rett i registeret, ikke her."
)
RELASJONSMERKNAD = (
    "Generert fra docs/datamodell/relasjoner.toml, dataflyt.toml og katalog.json av "
    "docs/verktoy/datamodell.py. Rett i registrene, ikke her."
)
HENDELSESMERKNAD = (
    "Generert fra docs/datamodell/hendelser.toml og backend-modellene av "
    "docs/verktoy/datamodell.py. Rett i katalogen eller modellene, ikke her."
)


def les(sti: pathlib.Path = REGISTER) -> dict:
    with sti.open("rb") as fil:
        return tomllib.load(fil)


def tekst(verdi) -> str:
    """Slår sammen linjeskift og mellomrom, slik at TOML kan brytes fritt."""
    if isinstance(verdi, list):
        return "; ".join(tekst(v) for v in verdi)
    return " ".join(str(verdi).split())


def ja_delvis_nei(typer: list[str], type_: str) -> str:
    if type_ not in typer:
        return "Nei"
    return "Ja" if len(typer) == 1 else "Delvis"


def funnregisteret(sti: pathlib.Path = HOVEDPLAN) -> str:
    """Avsnitt 4 i hovedplanen, der funn-ID-ene er ført."""
    plan = sti.read_text(encoding="utf-8")
    start = plan.index("## 4. Funnregister")
    return plan[start : plan.index("## 5. Arbeidspakker", start)]


def valider(register: dict, funntekst: str | None = None) -> list[str]:
    feil = []
    sett = set()
    for oppføring in register.get("tabell", []):
        navn = oppføring.get("navn", "<uten navn>")
        ukjente = set(oppføring) - set(TABELLFELT)
        mangler = [f for f in TABELLFELT if f not in oppføring]
        if ukjente:
            feil.append(f"{navn}: ukjente felt {sorted(ukjente)}")
        if mangler:
            feil.append(f"{navn}: mangler {mangler}")
            continue
        if navn in sett:
            feil.append(f"{navn}: står to ganger")
        sett.add(navn)
        if oppføring["lagring"] not in LAGRING:
            feil.append(f"{navn}: lagring {oppføring['lagring']!r} er ikke en av {LAGRING}")
        if oppføring["status"] not in STATUS:
            feil.append(f"{navn}: status {oppføring['status']!r} er ikke en av {STATUS}")
        typer = oppføring["type"]
        if not typer or not set(typer) <= set(TYPER):
            feil.append(f"{navn}: type {typer!r} må være en ikke-tom del av {sorted(TYPER)}")
        if not tekst(oppføring["kan_gjenoppbygges"]).startswith(GJENOPPBYGGING):
            feil.append(f"{navn}: kan_gjenoppbygges må begynne med en av {GJENOPPBYGGING}")
        for felt in ("beskrivelse", "lagres_fra", "fra_catenda", "personopplysninger", "belegg"):
            if not tekst(oppføring[felt]):
                feil.append(f"{navn}: {felt} er tomt; skriv «ikke kontrollert» hvis det er svaret")
        if oppføring["status"] == "i bruk" and not oppføring["skrivere"]:
            feil.append(f"{navn}: en tabell i bruk uten skriver er et funn, ikke en oppføring")
        for funn in oppføring["funn"]:
            if not FUNN_ID.match(funn):
                feil.append(f"{navn}: {funn!r} er ikke en funn-ID")
            elif funntekst is not None and not re.search(rf"\b{re.escape(funn)}\b", funntekst):
                feil.append(f"{navn}: {funn} står ikke i hovedplanens funnregister")
    for oppføring in register.get("utenfor", []):
        navn = oppføring.get("navn", "<uten navn>")
        if set(oppføring) != set(UTENFOR_FELT):
            feil.append(f"{navn}: feltene skal være {UTENFOR_FELT}")
        elif oppføring["lagring"] in LAGRING:
            feil.append(f"{navn}: står under «utenfor», men lagres i {oppføring['lagring']}")
    return feil


def avvik(register: dict, lagring: str, faktiske: set[str]) -> tuple[list[str], list[str]]:
    """Tabeller uten oppføring, og oppføringer uten tabell, for én lagring."""
    ført = {t["navn"] for t in register["tabell"] if t["lagring"] == lagring}
    return sorted(faktiske - ført), sorted(ført - faktiske)


def kilder() -> dict:
    """Relasjonsregisteret, dataflyten og katalogen, lest fra disk."""
    return {"relasjoner": rel.les(), "dataflyt": df.les(), "katalog": kat.les()}


def catenda_tekst(t: dict, k: dict) -> str:
    """Kolonne 6: vurderingen i tabellregisteret, og flytene som berører tabellen."""
    return f"{tekst(t['fra_catenda'])} {df.tabelltekst(k['dataflyt'], t['navn'])}"


def relasjonstekst(t: dict, k: dict) -> str:
    """Kolonne 7: nøkler og relasjoner, fra relasjonsregisteret og katalogen."""
    return rel.tabelltekst(k["relasjoner"], k["katalog"], t["navn"])


def tabellrader(register: dict, k: dict | None = None) -> list[list[str]]:
    k = k or kilder()
    rader = [IKT_KOLONNER + TILLEGGSKOLONNER]
    for t in register["tabell"]:
        rader.append(
            [
                t["navn"],
                tekst(t["beskrivelse"]),
                ja_delvis_nei(t["type"], "transaksjon"),
                tekst(t["lagres_fra"]),
                ja_delvis_nei(t["type"], "grunndata"),
                catenda_tekst(t, k),
                relasjonstekst(t, k),
                t["lagring"],
                ", ".join(t["type"]),
                t["status"],
                tekst(t["personopplysninger"]),
            ]
        )
    return rader


def utenfor_rader(register: dict) -> list[list[str]]:
    rader = [UTENFOR_KOLONNER]
    for u in register.get("utenfor", []):
        rader.append([tekst(u[f]) for f in UTENFOR_EXCEL])
    return rader


def om_rader() -> list[list[str]]:
    rader = [["Emne", "Forklaring"], ["Kilde", MERKNAD]]
    rader += [[f"Type: {navn}", forklaring] for navn, forklaring in TYPER.items()]
    rader.append(
        [
            "IKTs ja/nei-kolonner",
            "Ja når typen er den eneste, Delvis når den er én av flere, ellers Nei.",
        ]
    )
    rader.append(
        [
            "Ikke kontrollert",
            "Står der påstanden ikke er kontrollert. Registeret bygges i runder.",
        ]
    )
    rader.append(
        [
            "Detaljer",
            "Hvem som skriver hver tabell, belegg og funn står i docs/datamodell/tabeller.md i repoet.",
        ]
    )
    rader += [
        [
            "Relasjoner",
            (
                "Én rad per kobling mellom to tabeller. «Fremmednøkkel» håndheves av basen; "
                "«Uten fremmednøkkel» er en kobling koden bruker, som ingenting i basen håndhever. "
                "Nøklene og fremmednøklene er lest fra databasekatalogen."
            ),
        ],
        [
            "Dataflyt",
            (
                "Én rad per pil mellom en tabell og Catenda, med endepunktet og koden eller "
                "databasefunksjonen som gjør kallet. «Appen, lagres ikke» betyr at dataene "
                "brukes eller vises uten å bli lagret."
            ),
        ],
        [
            "Diagrammer",
            (
                "ER-diagrammene og flytdiagrammene står i docs/datamodell/relasjoner.md i repoet, "
                "generert fra de samme registrene."
            ),
        ],
        [
            "Hendelsestyper",
            (
                "Én rad per hendelsestype i tabellen hendelse. Innholdet ligger i kolonnen data "
                "og varierer med typen. Interne notater ligger i tabellen notat."
            ),
        ],
        [
            "Hendelsesfelt",
            (
                "Feltene i data per hendelsestype, generert fra modellene i koden. "
                "Delmodellene, som varselinformasjon og brev, står til slutt."
            ),
        ],
        [
            "Kilde for bestemmelsen",
            (
                "vedtak: et vedtak i hovedplanen eller en ADR. kode: koden oppgir bestemmelsen, "
                "ikke kontrollert mot standardteksten. oppdragsgiver: oppdragsgivers svar."
            ),
        ],
        [
            "Detaljer om hendelsene",
            "Innsendingsveier, forretningsregler, belegg og funn står i docs/datamodell/hendelser.md i repoet.",
        ],
    ]
    return rader


def ark(register: dict, katalog: dict, k: dict | None = None) -> dict[str, list[list[str]]]:
    k = k or kilder()
    return {
        "Tabeller": tabellrader(register, k),
        "Relasjoner": rel.relasjonsrader(k["relasjoner"], k["katalog"]),
        "Dataflyt": df.dataflytrader(k["dataflyt"], k["katalog"], register),
        "Utenfor databasene": utenfor_rader(register),
        "Hendelsestyper": hk.typerader(katalog),
        "Hendelsesfelt": hk.feltrader(katalog),
        "Om registeret": om_rader(),
    }


def _celle(verdi: str) -> str:
    return verdi.replace("|", "\\|")


def markdown(register: dict, k: dict | None = None) -> str:
    k = k or kilder()
    linjer = [
        "# Tabellregisteret",
        "",
        f"> {MERKNAD}",
        "",
        (
            "Hver tabell i appens databaser, og data som ligger utenfor dem. Feltene "
            "og typene er forklart i [oppdraget for spor M, avsnitt 5]"
            "(../prompt-datamodell-og-funksjonskart-2026-09-29.md#5-feltene-i-registeret). "
            "Registeret bygges i runder; «ikke kontrollert» betyr det det sier."
        ),
        "",
        "## Oversikt",
        "",
        "| Tabell | Lagring | Type | Status | Transaksjon | Grunndata |",
        "| --- | --- | --- | --- | --- | --- |",
    ]
    for t in register["tabell"]:
        linjer.append(
            f"| [`{t['navn']}`](#{t['navn']}) | {t['lagring']} | "
            f"{', '.join(t['type'])} | {t['status']} | "
            f"{ja_delvis_nei(t['type'], 'transaksjon')} | {ja_delvis_nei(t['type'], 'grunndata')} |"
        )
    linjer.append("")
    for lagring in LAGRING:
        linjer += [f"## {lagring}", ""]
        for t in register["tabell"]:
            if t["lagring"] != lagring:
                continue
            linjer += [
                f"### `{t['navn']}`",
                "",
                tekst(t["beskrivelse"]),
                "",
                "| Felt | Innhold |",
                "| --- | --- |",
                f"| Type | {', '.join(t['type'])} |",
                f"| Status | {t['status']} |",
                f"| Hvor fra appen lagres og redigeres | {_celle(tekst(t['lagres_fra']))} |",
                f"| Data fra Catenda | {_celle(catenda_tekst(t, k))} |",
                f"| Relasjoner | {_celle(relasjonstekst(t, k))} |",
                f"| Kan bygges opp igjen | {_celle(tekst(t['kan_gjenoppbygges']))} |",
                f"| Personopplysninger | {_celle(tekst(t['personopplysninger']))} |",
                f"| Skrives av | {_celle(tekst(t['skrivere']) or 'Ingen funnet')} |",
                f"| Funn | {', '.join(t['funn']) or '—'} |",
                f"| Belegg | {_celle(tekst(t['belegg']))} |",
                "",
            ]
    linjer += ["## Utenfor databasene", "", "| Navn | Lagring | Beskrivelse | Lagres fra | Belegg |", "| --- | --- | --- | --- | --- |"]
    for u in register.get("utenfor", []):
        linjer.append("| " + " | ".join(_celle(tekst(u[f])) for f in UTENFOR_FELT) + " |")
    linjer += ["", "## Typer og belegg", "", "| Kode | Betyr |", "| --- | --- |"]
    linjer += [f"| `{navn}` | {forklaring} |" for navn, forklaring in TYPER.items()]
    linjer += [f"| {kode} | {forklaring} |" for kode, forklaring in BELEGG.items()]
    return "\n".join(linjer) + "\n"


def _uten_tomme_til_slutt(rad) -> list[str]:
    verdier = ["" if c is None else str(c) for c in rad]
    while verdier and verdier[-1] == "":
        verdier.pop()
    return verdier


def excel_verdier(sti: pathlib.Path) -> dict[str, list[list[str]]]:
    from openpyxl import load_workbook

    bok = load_workbook(sti, read_only=True)
    try:
        return {
            arket.title: [_uten_tomme_til_slutt(rad) for rad in arket.iter_rows(values_only=True)]
            for arket in bok.worksheets
        }
    finally:
        bok.close()


def skriv_excel(register: dict, katalog: dict, sti: pathlib.Path, k: dict | None = None) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    bok = Workbook()
    bok.remove(bok.active)
    bredder = {
        "Tabeller": [28, 60, 16, 50, 16, 40, 50, 12, 18, 10, 36],
        "Relasjoner": [24, 30, 24, 24, 18, 16, 14, 70],
        "Dataflyt": [8, 30, 10, 40, 24, 24, 50, 60, 60],
        "Hendelsestyper": [28, 60, 12, 16, 18, 50, 40, 60, 10],
        "Hendelsesfelt": [30, 28, 30, 10, 60],
    }
    for tittel, rader in ark(register, katalog, k).items():
        arket = bok.create_sheet(tittel)
        for rad in rader:
            arket.append(rad)
        for celle in arket[1]:
            celle.font = Font(bold=True)
        for rad in arket.iter_rows():
            for celle in rad:
                celle.alignment = Alignment(wrap_text=True, vertical="top")
        arket.freeze_panes = "B2"
        standard = [30, 18, 70, 45, 50][: len(rader[0])]
        for i, bredde in enumerate(bredder.get(tittel, standard), start=1):
            arket.column_dimensions[arket.cell(row=1, column=i).column_letter].width = bredde
    bok.save(sti)


def kontroller_hendelser(katalog: dict) -> list[str]:
    feil = hk.valider(katalog, funnregisteret())
    uten_oppforing, uten_type = hk.avvik(katalog, hk.typene_i_koden())
    feil += [f"{t}: mangler oppføring i hendelser.toml" for t in uten_oppforing]
    feil += [f"{t}: finnes ikke i EventType" for t in uten_type]
    return feil + (hk.mot_koden(katalog) if not feil else [])


def katalogavvik(register: dict, skjema: dict) -> list[str]:
    """Tabellene i katalog.json mot tabellregisteret. Et avvik betyr at en av dem er utdatert."""
    feil = []
    for lag, lagring in (("postgresql", "PostgreSQL"), ("sqlite", "SQLite")):
        uten_oppforing, uten_tabell = avvik(register, lagring, set(skjema[lag]["tabeller"]))
        feil += [f"{t}: står i katalog.json, men ikke i tabeller.toml" for t in uten_oppforing]
        feil += [f"{t}: står i tabeller.toml, men ikke i katalog.json" for t in uten_tabell]
    return feil


def kontroller_relasjoner(register: dict, k: dict) -> list[str]:
    funntekst = funnregisteret()
    feil = katalogavvik(register, k["katalog"])
    feil += rel.valider(k["relasjoner"], k["katalog"], funntekst)
    feil += df.valider(k["dataflyt"], k["katalog"], register, funntekst)
    per_tabell = df.flyter_per_tabell(k["dataflyt"])
    for t in register["tabell"]:
        inn = per_tabell.get(t["navn"], {}).get("fra")
        if inn and tekst(t["fra_catenda"]).startswith("Nei"):
            feil.append(f"{t['navn']}: fra_catenda sier «Nei», men dataflyten henter fra Catenda i {', '.join(inn)}")
    steder = df.kallsteder(frozenset(k["dataflyt"].get("autentisering", [])))
    feil += [f"{sted}: kaller Catenda uten pil i dataflyt.toml" for sted in df.udekkede_kallsteder(k["dataflyt"], steder)]
    return feil


def relasjoner_markdown(register: dict, k: dict | None = None) -> str:
    k = k or kilder()
    relasjoner, flyt, skjema = k["relasjoner"], k["dataflyt"], k["katalog"]
    fk = rel.fremmednøkler(skjema)
    linjer = [
        "# Relasjoner og dataflyt",
        "",
        f"> {RELASJONSMERKNAD}",
        "",
        (
            "Hvordan tabellene henger sammen, og hvilke av dem som henter fra eller sender til "
            "Catenda. Nøklene og fremmednøklene er lest fra databasekatalogen "
            "([`katalog.json`](katalog.json)); koblingene uten fremmednøkkel og dataflyten er "
            "lest ut av koden. Tabellene er beskrevet i [tabellregisteret](tabeller.md). "
            "Registeret bygges i runder; «ikke kontrollert» betyr det det sier."
        ),
        "",
        "## ER-diagram: PostgreSQL",
        "",
        (
            "Heltrukken linje er en fremmednøkkel. Stiplet linje er en kobling uten "
            "fremmednøkkel, som basen ikke håndhever. `auth_users` er tabellen `auth.users` "
            "i Supabase-plattformen."
        ),
        "",
        "```mermaid",
        rel.er_diagram(relasjoner, skjema, "PostgreSQL"),
        "```",
        "",
        "## ER-diagram: SQLite",
        "",
        (
            f"De {len(skjema['sqlite']['tabeller'])} tabellene i fila `BH_APPROVAL_DB`, og "
            "tabellene i PostgreSQL de viser til. "
            "Ingen av koblingene kan håndheves av en base, fordi de går mellom to lagre."
        ),
        "",
        "```mermaid",
        rel.er_diagram(relasjoner, skjema, "SQLite"),
        "```",
        "",
        "## Relasjonene",
        "",
        "| Fra | Til | Art | Antall | Beskrivelse | Funn | Belegg |",
        "| --- | --- | --- | --- | --- | --- | --- |",
    ]
    for r in relasjoner["relasjon"]:
        art = (
            f"fremmednøkkel, ON DELETE {fk[(r['fra'], r['til'])]}"
            if r["art"] == "fremmednøkkel"
            else "uten fremmednøkkel"
        )
        linjer.append(
            f"| `{rel._side(r, 'fra')}` | `{rel._side(r, 'til')}` | {art} | {r['antall']} | "
            f"{_celle(tekst(r['beskrivelse']))} | {', '.join(r['funn']) or '—'} | {_celle(tekst(r['belegg']))} |"
        )
    linjer += [
        "",
        "## Nøkkelkolonner uten relasjon",
        "",
        (
            "Kolonner som ser ut som en nøkkel eller en personreferanse, men som ikke viser til en "
            "tabell i appen. Testen krever en begrunnelse for hver."
        ),
        "",
        "| Kolonne | Begrunnelse |",
        "| --- | --- |",
    ]
    linjer += [
        f"| `{u['kolonne']}` | {_celle(tekst(u['begrunnelse']))} |" for u in relasjoner["uten_relasjon"]
    ]
    linjer += [
        "",
        "## Dataflyt mot Catenda",
        "",
        (
            "Hver pil er nummerert med flyten og steget, og viser siste ledd i kallkjeden: "
            "koden som gjør kallet, eller databasefunksjonen som skriver. Hele kjeden og "
            "endepunktet står i flytene under. «Appen, lagres ikke» betyr at dataene brukes "
            "eller vises uten å bli lagret."
        ),
        "",
        "### Inn fra Catenda",
        "",
        "```mermaid",
        df.flytdiagram(flyt, skjema, register, "inn"),
        "```",
        "",
        "### Ut til Catenda",
        "",
        "```mermaid",
        df.flytdiagram(flyt, skjema, register, "ut"),
        "```",
        "",
        "### Drift",
        "",
        (
            "Flytene bare driftsskriptene utløser. De er den eneste veien prosjekter og "
            "kontraktsteam kommer inn i basen. Medlemssynkroniseringen (C04) står i "
            "diagrammet over, fordi også innloggingen og tilgangskontrollen utløser den."
        ),
        "",
        "```mermaid",
        df.flytdiagram(flyt, skjema, register, "drift"),
        "```",
        "",
        "### Tabellene og Catenda",
        "",
        "| Tabell | Får data fra Catenda i | Sender data til Catenda i | Kvittering for kall i |",
        "| --- | --- | --- | --- |",
    ]
    per_tabell = df.flyter_per_tabell(flyt)
    for t in register["tabell"]:
        flyter = per_tabell.get(t["navn"], {"fra": [], "til": [], "kvittering": []})
        linjer.append(
            f"| `{t['navn']}` | "
            + " | ".join(", ".join(flyter[k]) or "—" for k in ("fra", "til", "kvittering"))
            + " |"
        )
    linjer += ["", "## Flytene", ""]
    linjer += df.flytseksjoner(flyt, skjema, register)
    linjer += [
        "## Kall til Catenda utenfor dataflyten",
        "",
        (
            "Steder i backend som kaller Catenda-klienten, men ikke er en del av appens dataflyt. "
            f"Autentiseringskallene ({', '.join(f'`{a}`' for a in flyt['autentisering'])}) henter "
            "bare et token og er ikke tatt med."
        ),
        "",
        "| Fil | Funksjon | Begrunnelse |",
        "| --- | --- | --- |",
    ]
    linjer += [
        f"| [`{u['fil'].removeprefix('backend/')}`](../../{u['fil']}) | "
        f"{('`' + u['funksjon'] + '`') if 'funksjon' in u else 'hele fila'} | {_celle(tekst(u['begrunnelse']))} |"
        for u in flyt["uten_flyt"]
    ]
    return "\n".join(linjer) + "\n"


def main() -> int:
    register = les()
    katalog = hk.les()
    k = kilder()
    feil = valider(register, funnregisteret()) + kontroller_hendelser(katalog)
    feil += kontroller_relasjoner(register, k)
    if feil:
        print("Registeret eller katalogen er ikke gyldig:", *feil, sep="\n  ")
        return 1
    for sti, ny in (
        (MARKDOWN, markdown(register, k)),
        (RELASJONER, relasjoner_markdown(register, k)),
        (hk.MARKDOWN, hk.markdown(katalog, HENDELSESMERKNAD)),
    ):
        if not sti.exists() or sti.read_text(encoding="utf-8") != ny:
            sti.write_text(ny, encoding="utf-8")
            print(f"skrev {sti.relative_to(ROT)}")
    if not EXCEL.exists() or excel_verdier(EXCEL) != ark(register, katalog, k):
        skriv_excel(register, katalog, EXCEL, k)
        print(f"skrev {EXCEL.relative_to(ROT)}")
    print(
        f"{len(register['tabell'])} tabeller, {len(register.get('utenfor', []))} utenfor databasene, "
        f"{len(katalog['hendelse'])} hendelsestyper, {len(k['relasjoner']['relasjon'])} relasjoner, "
        f"{len(df.piler(k['dataflyt']))} piler i dataflyten"
    )
    return 0


if __name__ == "__main__":
    sys.exit(main())
